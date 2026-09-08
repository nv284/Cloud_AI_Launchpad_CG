from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()
    
    # Helper function to add a slide with title and bullet points
    def add_slide(title, content_list):
        slide_layout = prs.slide_layouts[1] # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        title_shape.text = title
        
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        
        for i, point in enumerate(content_list):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = point
            p.font.size = Pt(18)
            p.level = 0

    # Slide 1: Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Python for Data Handling & AI Applications"
    slide.placeholders[1].text = "Modules 10-19: From Data Analysis to Generative AI\nComprehensive Course Presentation"

    # Slide 2: Module 10
    add_slide("Module 10: Python for Data Handling", [
        "• NumPy Introduction: The foundational library for numerical computing in Python.",
        "• Arrays & Dimensions: Creating 1D, 2D, and N-dimensional arrays (ndarrays).",
        "• Indexing & Slicing: Efficiently accessing and modifying array elements.",
        "• Vectorized Operations: Performing math on entire arrays without slow for-loops (broadcasting)."
    ])

    # Slide 3: Module 11
    add_slide("Module 11: Data Analysis with Pandas", [
        "• Core Structures: Series (1D) and DataFrames (2D) for structured data.",
        "• Data Ingestion: Reading data using read_csv() and read_json().",
        "• Data Manipulation: Filtering rows, sorting values, and selecting columns.",
        "• Data Cleaning: Handling missing data (dropna, fillna) and basic transformations (apply, groupby)."
    ])

    # Slide 4: Module 12
    add_slide("Module 12: Data Visualization", [
        "• Matplotlib Basics: Importing pyplot and creating figures/axes.",
        "• Plots & Charts: Creating line, scatter, bar, and histogram plots.",
        "• Customization: Adding titles, x/y labels, legends, and adjusting colors.",
        "• Interpretation: Analyzing visualizations to identify trends, outliers, and correlations."
    ])

    # Slide 5: Module 13
    add_slide("Module 13: Working with APIs & JSON", [
        "• HTTP Basics: Understanding GET, POST, PUT, DELETE methods and status codes.",
        "• REST APIs: Architectural style for web services; stateless communication.",
        "• The 'requests' Library: Making HTTP calls and handling responses in Python.",
        "• JSON Parsing: Using the json module to parse API responses and consume AI services."
    ])

    # Slide 6: Module 14
    add_slide("Module 14: Python Environment & AI Libraries", [
        "• Virtual Environments: Using 'venv' to isolate project dependencies.",
        "• Dependency Management: Installing packages via 'pip' and freezing them.",
        "• Requirements Files: Using requirements.txt to ensure reproducible environments.",
        "• AI/ML Libraries: Introduction to Scikit-learn, TensorFlow, PyTorch, and HuggingFace."
    ])

    # Slide 7: Module 15
    add_slide("Module 15: Python for AI Applications", [
        "• The AI Workflow: A structured pipeline for building AI apps.",
        "• 1. Data Ingestion: Gathering raw data from files or APIs.",
        "• 2. Preprocessing: Cleaning, normalizing, and formatting data for models.",
        "• 3. Model/API Execution: Passing data to an ML model or GenAI API.",
        "• 4. Response & Integration: Parsing the output and building the final application UI/Logic."
    ])

    # Slide 8: Module 16
    add_slide("Module 16: Introduction to Machine Learning", [
        "• AI vs ML vs Deep Learning: AI (broad), ML (learning from data), DL (neural networks).",
        "• Supervised Learning: Training on labeled data (Features X → Labels y).",
        "• Unsupervised Learning: Finding hidden patterns in unlabeled data (clustering).",
        "• Basic Workflow: Data split (train/test) → Model selection → Training → Evaluation."
    ])

    # Slide 9: Module 17
    add_slide("Module 17: Generative AI with Python", [
        "• LLM Concepts: Understanding Large Language Models and tokenization.",
        "• Prompt Engineering: Crafting effective system and user prompts.",
        "• API Interaction: Sending prompts to OpenAI/Anthropic via Python 'requests'.",
        "• Processing Responses: Handling streaming text, JSON modes, and extracting structured data."
    ])

    # Slide 10: Module 18
    add_slide("Module 18: Mini Project - Python AI Assistant", [
        "• Project Goal: Build a functional CLI/Python AI Assistant.",
        "• Architecture: Combining data processing, API calls, and modular functions.",
        "• Implementation: Fetching data → Processing with Pandas → Querying GenAI.",
        "• Output: Generating dynamic, context-aware responses for the user."
    ])

    # Slide 11: Module 19
    add_slide("Module 19: Assessment & Wrap-up", [
        "• Coding Challenge: Clean a messy dataset and perform exploratory analysis.",
        "• Troubleshooting Exercise: Debug a broken API call or ML pipeline.",
        "• Knowledge Check: Quiz on AI concepts, Pandas syntax, and HTTP methods.",
        "• Next Steps: Best practices for production code and advanced learning paths."
    ])

    # Slide 12: Integration
    add_slide("Connecting the Dots: The Data-to-AI Pipeline", [
        "• Data Foundation (Mod 10-12): NumPy and Pandas clean the data; Matplotlib visualizes it.",
        "• Infrastructure (Mod 13-14): Virtual environments keep it stable; APIs fetch external data.",
        "• Intelligence (Mod 15-17): The cleaned data feeds into ML models (predictions) or LLMs (generation).",
        "• Application (Mod 18-19): All components are wrapped into a robust, tested AI Assistant."
    ])

    # Slide 13: Best Practices
    add_slide("Best Practices for Python AI Development", [
        "• Code Modularity: Use functions and classes to separate data, logic, and API calls.",
        "• Error Handling: Always use try/except blocks for API calls and file I/O.",
        "• Security: Never hardcode API keys; use environment variables (.env).",
        "• Performance: Use vectorized NumPy/Pandas operations instead of iterative loops."
    ])

    # Slide 14: Summary
    add_slide("Course Summary", [
        "• Mastered Python data manipulation with NumPy and Pandas.",
        "• Created insightful visualizations using Matplotlib.",
        "• Integrated external data and AI services via REST APIs and JSON.",
        "• Built a complete pipeline from data preprocessing to Machine Learning and Generative AI applications."
    ])

    # Slide 15: Case Study (Last Slide)
    add_slide("Case Study: 'SmartRetail AI' End-to-End Assistant", [
        "• Scenario: Build an AI assistant that analyzes retail sales, predicts trends, and summarizes insights.",
        "• Data & Viz (Mod 10-12): Use Pandas to clean sales CSVs; NumPy for math; Matplotlib to plot monthly revenue.",
        "• APIs & Env (Mod 13-14): Use a venv; fetch live currency exchange rates via a REST API using 'requests'.",
        "• AI & ML (Mod 15-17): Preprocess data. Use Scikit-learn (ML) to predict next month's sales. Send top-selling items to an LLM API (GenAI) to generate an executive summary.",
        "• Project & Assessment (Mod 18-19): Wrap in a CLI Assistant. Implement error handling for API rate limits (troubleshooting) and modularize code."
    ])

    prs.save('Python_Data_and_AI_Course.pptx')
    print("Presentation saved successfully as 'Python_Data_and_AI_Course.pptx'")

if __name__ == "__main__":
    create_presentation()