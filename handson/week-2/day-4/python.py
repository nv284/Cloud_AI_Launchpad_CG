from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()

    # Helper function to add a slide with title and bullet points
    def add_slide(title, content_list):
        slide_layout = prs.slide_layouts[1] # Bullet layout
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        title_shape.text = title
        
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        
        for i, item in enumerate(content_list):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            
            # Handle sub-bullets
            if item.startswith("  - "):
                p.text = item.replace("  - ", "")
                p.level = 1
                p.font.size = Pt(16)
            else:
                p.text = item
                p.level = 0
                p.font.size = Pt(18)
                p.font.bold = True

    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Python for AI Integration & Automation"
    slide.placeholders[1].text = "Advanced Topics for Enterprise Professionals\nSeptember 2026"

    # Slide 2: Adv LLM 1
    add_slide("1. Advanced LLM Integration (Core)", [
        "LLM API Architecture: REST/gRPC endpoints, stateless request/response models.",
        "Python SDK Integration: Using official SDKs (OpenAI, Anthropic) and unified routers (LiteLLM).",
        "Structured Output: Enforcing Pydantic models and JSON modes for reliable data extraction.",
        "Function/Tool Calling: Defining schemas for LLMs to trigger external Python functions.",
        "Streaming Responses: Implementing Server-Sent Events (SSE) and async generators for real-time UI updates."
    ])

    # Slide 3: Adv LLM 2
    add_slide("1. Advanced LLM Integration (Resilience)", [
        "Token & Context Management: Sliding windows, context compression, and summarization techniques.",
        "Retry & Timeout Handling: Implementing exponential backoff (Tenacity) for transient API failures.",
        "Model Selection & Fallbacks: Semantic routing to route simple tasks to small models and complex tasks to frontier models.",
        "Cost Optimization: Caching identical prompts and batching asynchronous requests."
    ])

    # Slide 4: AI Automation Arch 1
    add_slide("2. AI-Powered Automation Architecture (Foundations)", [
        "AI vs Traditional Automation: Probabilistic (LLM) vs Deterministic (Rules) execution paradigms.",
        "API → AI → Action Workflow: Ingesting data via API, reasoning via LLM, and executing state changes.",
        "Event-Driven Automation: Triggering AI workflows via webhooks, Kafka, or message queues.",
        "Human-in-the-Loop (HITL): Pausing execution for human approval when AI confidence scores fall below thresholds."
    ])

    # Slide 5: AI Automation Arch 2
    add_slide("2. AI-Powered Automation Architecture (Execution)", [
        "Decision-Making with LLMs: Utilizing Chain-of-Thought (CoT) for complex logical routing.",
        "Trigger → Process → Validate → Execute Pattern: Ensuring AI outputs are validated against business rules before execution.",
        "Multi-Step Workflows: Chaining multiple AI and deterministic steps for complex enterprise processes.",
        "Enterprise Use Cases: Automated invoice processing, IT ticket routing, and dynamic compliance reporting."
    ])

    # Slide 6: RAG 1
    add_slide("3. RAG-Based AI Applications (Foundations)", [
        "RAG Architecture: Grounding LLM responses in proprietary, real-time enterprise data.",
        "Document Ingestion Pipeline: Parsing PDFs, HTML, and DBs while preserving layout and tables.",
        "Chunking Strategies: Recursive character splitting for context preservation; semantic chunking for meaning boundaries.",
        "Embeddings: Transforming text into high-dimensional vectors (Dense vs. Sparse representations).",
        "Vector Databases: Storing and querying vectors efficiently (Pinecone, Milvus, pgvector)."
    ])

    # Slide 7: RAG 2
    add_slide("3. RAG-Based AI Applications (Advanced)", [
        "Semantic Search & Metadata Filtering: Combining vector similarity with strict keyword/metadata filters (Hybrid Search).",
        "Retrieval → Context → LLM Pipeline: Prompt engineering to synthesize retrieved chunks accurately.",
        "RAG Evaluation: Using frameworks like RAGAS to measure context precision, faithfulness, and answer relevance.",
        "Common Failure Points: 'Lost in the middle' phenomenon, chunk boundary splits, and hallucination.",
        "Python Implementation: Leveraging LlamaIndex and LangChain for modular RAG pipelines."
    ])

    # Slide 8: Agents 1
    add_slide("4. AI Agents & Tool Calling (Core)", [
        "Agent vs Chatbot vs Workflow: Agents possess autonomy, memory, and tool-use; chatbots are conversational; workflows are static.",
        "Agent Architecture: Perception (inputs), Brain (LLM reasoning), Action (tool execution).",
        "The Agent Loop: Observe environment → Think (plan) → Act (execute tool) → Observe results.",
        "Planning & Execution: Implementing ReAct (Reason + Act) and Plan-and-Solve frameworks.",
        "Tool/Function Calling: Securely mapping LLM intents to authenticated Python API/DB calls."
    ])

    # Slide 9: Agents 2
    add_slide("4. AI Agents & Tool Calling (Advanced)", [
        "Agent Memory: Short-term (context window) vs Long-term (vector DB) memory for persistent state.",
        "Multi-Agent Architecture: Orchestrating specialized agents (e.g., Researcher, Coder, Reviewer) via CrewAI or AutoGen.",
        "LangGraph Concepts: Building stateful, multi-actor applications with cycles and conditional branching.",
        "MCP (Model Context Protocol): Standardizing how agents connect to external data sources and tools securely."
    ])

    # Slide 10: Prod Ready 1
    add_slide("5. Building Production-Ready AI (Backend)", [
        "FastAPI for AI Services: High-performance, async Python framework for exposing AI endpoints.",
        "Async Processing: Offloading long-running LLM tasks to Celery/Redis to prevent API timeouts.",
        "Authentication & Security: OAuth2, API key management, and strict PII masking before LLM inference.",
        "Logging & Observability: Tracing LLM calls, token usage, and latency using Langfuse, LangSmith, or OpenTelemetry."
    ])

    # Slide 11: Prod Ready 2
    add_slide("5. Building Production-Ready AI (Ops & Deploy)", [
        "Retry & Recovery: Implementing circuit breakers to prevent cascading failures during LLM outages.",
        "Rate Limiting & Caching: Protecting APIs and using Semantic Caching (Redis) to reduce redundant LLM costs.",
        "Environment & Secrets: Managing keys securely via HashiCorp Vault or AWS Secrets Manager.",
        "Docker Deployment: Containerizing AI services with optimized base images and GPU passthrough.",
        "Cost & Performance: Prompt compression, model routing, and batch processing for enterprise scale."
    ])

    # Slide 12: Case Study
    add_slide("6. Case Study: Intelligent Customer Support", [
        "Python Components: FastAPI (gateway), LangGraph (agent orchestration), pgvector (RAG).",
        "LLM vs Deterministic: LLM for intent/summarization; Deterministic for routing, refunds, and SLA checks.",
        "RAG Implementation: Ingesting Zendesk/Confluence docs; hybrid search for accurate policy retrieval.",
        "Agent Tool Invocation: Agent calls CRM API to fetch user history, then executes ticket update.",
        "Failures & Security: Fallback to human if confidence < 80%; PII redaction before LLM context injection.",
        "Monitoring: Tracking token costs per resolution, latency percentiles, and customer satisfaction (CSAT)."
    ])

    prs.save('Python_AI_Integration_Advanced_Topics.pptx')
    print("Presentation saved successfully as 'Python_AI_Integration_Advanced_Topics.pptx'")

if __name__ == "__main__":
    create_presentation()