from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    prs = Presentation()

    # Helper function to add a content slide
    def add_slide(title_text, content_text, notes_text=""):
        slide_layout = prs.slide_layouts[1] # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = title_text
        
        content = slide.placeholders[1]
        content.text = content_text
        
        if notes_text:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = notes_text

    # Slide 1: Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "AI for Developers: APIs to Enterprise Productivity"
    slide.placeholders[1].text = "A Comprehensive Guide for Freshers\nPresented by: [Your Name/Company]\nDate: September 11, 2026"

    # Slide 2: AI API Fundamentals
    add_slide(
        "1. AI API Fundamentals",
        "• AI API: A bridge allowing your application to communicate with an AI model over the internet.\n"
        "• LLM (Large Language Model): The 'brain' (e.g., GPT, Claude) that processes and generates text.\n"
        "• Prompt: The specific input, context, or instruction sent to the AI.\n"
        "• Token: The basic unit of text (words/characters) the AI reads/generates; dictates cost and limits.\n"
        "• Request & Response: The structured data payload sent to the API, and the AI's structured reply.",
        "Speaker Notes: Spend time explaining that an API is just a messenger. The LLM is the brain, but it has no eyes or ears. The API provides the interface. Emphasize that tokens are how we pay for AI—it's not per word, but per chunk of characters."
    )

    # Slide 3: Calling an AI API
    add_slide(
        "2. Calling an AI API from an Application",
        "• API Key: Your unique secret password for authentication (keep it secure!).\n"
        "• Endpoint: The specific URL where the request is sent (e.g., api.openai.com/v1/chat/completions).\n"
        "• Authentication: Proving your identity via HTTP headers (e.g., Authorization: Bearer <KEY>).\n"
        "• JSON: The standard data format for sending the prompt and receiving the response.\n"
        "• Workflow: App formats JSON -> Sends HTTP POST -> API processes -> Returns JSON response.",
        "Speaker Notes: Highlight security here. Freshers often hardcode API keys. Teach them to use environment variables (.env files) immediately. Show a quick mental model of what the JSON request and response look like."
    )

    # Slide 4: Prompt → AI Model → Response Flow
    add_slide(
        "3. The AI Execution Flow",
        "• Step 1: User Input - User interacts with your app (typing a question).\n"
        "• Step 2: Application Logic - Your code formats the input into a Prompt and adds the API Key.\n"
        "• Step 3: API Request - App sends an HTTP POST request to the AI provider's endpoint.\n"
        "• Step 4: AI Processing - The LLM processes tokens, calculates probabilities, and generates text.\n"
        "• Step 5: Response & Display - API returns text to your app, which parses and displays it to the user.",
        "Speaker Notes: Walk through this like a story. Use the analogy of ordering at a restaurant: User (Customer) -> App (Waiter) -> API (Kitchen Door) -> LLM (Chef) -> Response (Food). "
    )

    # Slide 5: Structured AI Responses
    add_slide(
        "4. Structured AI Responses",
        "• The Problem: Plain text is hard for applications to parse and use programmatically.\n"
        "• The Solution: Instruct the AI to return data in structured formats like JSON or XML.\n"
        "• How it Works: Use system prompts (e.g., 'Return answer strictly as JSON with keys: name, email').\n"
        "• Benefits: Enables seamless integration with databases, UI components, and automated workflows.\n"
        "• Pro-Tip: Use modern API features like 'JSON Mode' or 'Function Calling' to guarantee valid output.",
        "Speaker Notes: This is a crucial concept for developers. If the AI returns plain text, you have to write fragile regex to parse it. If it returns JSON, you can directly map it to your database models. Mention 'Function Calling' as the industry standard for this."
    )

    # Slide 6: Developer Copilots
    add_slide(
        "5. Developer Copilots",
        "• What are they? AI tools integrated directly into your IDE (e.g., GitHub Copilot, Cursor, Codeium).\n"
        "• How they work: They analyze your current file, open tabs, and cursor position for context.\n"
        "• Core Features:\n"
        "   - Real-time code completion (ghost text).\n"
        "   - Chat interfaces for asking coding questions.\n"
        "   - Inline commands (e.g., 'Refactor this', 'Add comments').\n"
        "• Mindset Shift: Treat the Copilot as a junior pair programmer—review its work, don't blindly accept it.",
        "Speaker Notes: Emphasize the 'Mindset Shift'. Copilots are incredibly fast but can introduce subtle bugs or security vulnerabilities. The developer is still the pilot; the AI is just the co-pilot."
    )

    # Slide 7: AI-Assisted Developer Productivity
    add_slide(
        "6. Boosting Developer Productivity",
        "• Code Generation: Writing boilerplate, CRUD operations, or complex algorithms from natural language.\n"
        "• Explanation: Breaking down legacy or complex code into understandable, step-by-step logic.\n"
        "• Debugging: Identifying bugs, explaining stack traces, and suggesting targeted fixes.\n"
        "• Refactoring: Improving code readability, performance, and adherence to best practices.\n"
        "• Test Generation: Automatically creating unit tests and edge-case scenarios.\n"
        "• Documentation: Generating docstrings, READMEs, and API documentation instantly.",
        "Speaker Notes: Give practical examples. For instance, 'Instead of spending 2 hours writing unit tests for a utility function, prompt the AI to write them, then spend 15 minutes reviewing and tweaking them.'"
    )

    # Slide 8: Enterprise Productivity Use Cases
    add_slide(
        "7. Enterprise Productivity Use Cases",
        "• Communication: Summarizing long email threads, drafting replies, generating meeting notes.\n"
        "• Knowledge Management: Semantic search across company wikis (finding answers, not just documents).\n"
        "• Customer/IT Support: AI ticket assistance, auto-categorizing tickets, suggesting resolutions.\n"
        "• Data Analysis: Extracting insights from unstructured data (e.g., summarizing 50 vendor contracts).\n"
        "• Impact: Shifts employee focus from repetitive tasks to high-value, strategic work.",
        "Speaker Notes: Connect this to business value. AI isn't just for writing code; it's for accelerating the entire business. Mention RAG (Retrieval-Augmented Generation) briefly as the tech behind enterprise knowledge search."
    )

    # Slide 9: Responsible Enterprise AI Usage
    add_slide(
        "8. Responsible AI & Best Practices",
        "• Data Privacy: Never input PII, passwords, or proprietary code into public AI models.\n"
        "• Security: Keep API keys out of source code (use environment variables/secrets managers).\n"
        "• Hallucinations: AI can confidently generate false information. Always validate critical outputs.\n"
        "• Human in the Loop: AI is an assistant, not the final decision-maker. Humans must review and approve.\n"
        "• Access Control: Implement Role-Based Access Control (RBAC) so users only see authorized data.",
        "Speaker Notes: End on a strong note about responsibility. A single leaked API key or a hallucinated piece of code pushed to production can cause major enterprise incidents. 'Trust, but verify' is the golden rule."
    )

    prs.save('AI_Developer_Fundamentals_Presentation.pptx')
    print("Presentation saved successfully as 'AI_Developer_Fundamentals_Presentation.pptx'")

if __name__ == "__main__":
    create_presentation()